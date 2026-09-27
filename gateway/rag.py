import re

import httpx
import numpy as np
from django.conf import settings
from pgvector.django import CosineDistance

from .models import Chunk, Document

SUPPORTED = (".pdf", ".txt", ".md", ".docx")

# ── Embedding (calls the embeddings process, which holds the model) ──────
_sync: httpx.Client | None = None
_async: httpx.AsyncClient | None = None
EMBED_PATH = "/api/v1/internal/embed/"


def _headers() -> dict:
    return {"Authorization": f"Bearer {settings.INTERNAL_TOKEN}"}


def embed(texts: list[str], kind: str = "passage") -> np.ndarray:
    """Sync, for Celery. Returns an L2-normalized float32 matrix."""
    global _sync
    if _sync is None:
        _sync = httpx.Client(base_url=settings.EMBEDDER_URL, timeout=300, trust_env=False)
    r = _sync.post(EMBED_PATH, json={"texts": texts, "kind": kind}, headers=_headers())
    r.raise_for_status()
    return np.array(r.json()["vectors"], dtype=np.float32)


async def aembed(texts: list[str], kind: str = "query") -> np.ndarray:
    """Async, for chat requests."""
    global _async
    if _async is None:
        _async = httpx.AsyncClient(base_url=settings.EMBEDDER_URL, timeout=30, trust_env=False)
    r = await _async.post(EMBED_PATH, json={"texts": texts, "kind": kind}, headers=_headers())
    r.raise_for_status()
    return np.array(r.json()["vectors"], dtype=np.float32)


# ── Streaming extraction + chunking (used by the Celery task) ────────────
def iter_pages(path: str, filename: str):
    """Yield (page_number | None, text) one page at a time, so memory stays flat."""
    name = filename.lower()
    if name.endswith(".pdf"):
        import pymupdf  # older PyMuPDF versions: import fitz
        with pymupdf.open(path) as pdf:
            for i, page in enumerate(pdf, start=1):
                yield i, page.get_text()
    elif name.endswith((".txt", ".md")):
        with open(path, "r", encoding="utf-8", errors="ignore") as f:
            yield None, f.read()
    elif name.endswith(".docx"):
        import docx  # pip install python-docx
        d = docx.Document(path)
        parts = [p.text for p in d.paragraphs]
        for t in d.tables:
            for row in t.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        yield None, "\n".join(parts)
    else:
        raise ValueError("Only PDF, TXT, MD and DOCX files are supported")


def chunk_text(text: str) -> list[str]:
    size, overlap = settings.RAG_CHUNK_SIZE, settings.RAG_CHUNK_OVERLAP
    text = re.sub(r"\s+", " ", text).strip()
    chunks, i = [], 0
    while i < len(text):
        end = min(i + size, len(text))
        if end < len(text):
            cut = text.rfind(" ", i + size // 2, end)  # don't cut mid-word
            if cut != -1:
                end = cut
        chunks.append(text[i:end].strip())
        if end >= len(text):
            break
        i = max(end - overlap, i + 1)
    return [c for c in chunks if c]


def iter_chunks(path: str, filename: str):
    for page, text in iter_pages(path, filename):
        for c in chunk_text(text):
            yield page, c


def batched(iterable, n: int):
    batch = []
    for x in iterable:
        batch.append(x)
        if len(batch) == n:
            yield batch
            batch = []
    if batch:
        yield batch


# ── Retrieval ────────────────────────────────────────────────────────────
async def retrieve(key_id: int, question: str) -> list[str]:
    ready = Document.objects.filter(api_key_id=key_id, status=Document.Status.READY)
    if not await ready.aexists():          # no documents: skip embedding entirely
        return []

    q = (await aembed([question], "query"))[0]
    rows = [
        r async for r in
        Chunk.objects.filter(api_key_id=key_id, document__status=Document.Status.READY)
        .annotate(dist=CosineDistance("embedding", q))
        .order_by("dist")
        .values_list("document__filename", "page", "text", "dist")[: settings.RAG_TOP_K]
    ]
    rows.sort(key=lambda r: r[3])          # relaxed_order scans can be slightly out of order
    max_dist = 1 - settings.RAG_MIN_SCORE  # cosine distance = 1 - similarity
    return [
        f"[{name}, p.{page}] {text}" if page else f"[{name}] {text}"
        for name, page, text, dist in rows if dist <= max_dist
    ]